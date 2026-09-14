"""
report_generator.py
--------------------
Turns the list of ValidationResult objects into a single Excel workbook:
a Summary tab plus one tab per decision tier, mirroring how an AP team
would actually want to triage a batch.
"""

from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TIER_COLORS = {
    "AUTO_APPROVED": "1F7A3D",
    "NEEDS_REVIEW": "B7791F",
    "REJECTED": "B91C1C",
}

COLUMNS = [
    ("Invoice File", 16),
    ("Invoice #", 14),
    ("Vendor", 26),
    ("Entity", 26),
    ("PO Reference", 14),
    ("Currency", 10),
    ("Total Amount", 14),
    ("Flags", 60),
]


def _write_sheet(ws, results):
    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, r in enumerate(results, start=2):
        d = r.extracted
        flags_text = "; ".join(f"[{f['code']}] {f['detail']}" for f in r.flags) or "—"
        row = [
            r.invoice_id,
            d.get("invoice_number"),
            d.get("vendor"),
            d.get("bill_to_entity"),
            d.get("po_reference") or "—",
            d.get("currency"),
            d.get("total_amount"),
            flags_text,
        ]
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == 8:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 30 if r.flags else 15

    ws.freeze_panes = "A2"


def generate_report(results, output_path):
    tiers = {"AUTO_APPROVED": [], "NEEDS_REVIEW": [], "REJECTED": []}
    for r in results:
        tiers[r.decision].append(r)

    wb = Workbook()

    # --- Summary tab ---
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = f"Invoice Batch Report — {date.today().isoformat()}"
    summary["A1"].font = Font(bold=True, size=14)

    summary["A3"] = "Tier"
    summary["B3"] = "Count"
    summary["C3"] = "Total Value"
    for c in ("A3", "B3", "C3"):
        summary[c].font = HEADER_FONT
        summary[c].fill = HEADER_FILL

    row = 4
    for tier, label in [
        ("AUTO_APPROVED", "Auto-Approved"),
        ("NEEDS_REVIEW", "Needs Review"),
        ("REJECTED", "Rejected (duplicate)"),
    ]:
        items = tiers[tier]
        total_value = sum(r.extracted.get("total_amount") or 0 for r in items)
        summary.cell(row=row, column=1, value=label)
        summary.cell(row=row, column=2, value=len(items))
        summary.cell(row=row, column=3, value=round(total_value, 2))
        row += 1

    row += 1
    summary.cell(row=row, column=1, value="Flag breakdown (Needs Review + Rejected)").font = Font(bold=True)
    row += 1
    flag_counts = {}
    for r in tiers["NEEDS_REVIEW"] + tiers["REJECTED"]:
        for f in r.flags:
            flag_counts[f["code"]] = flag_counts.get(f["code"], 0) + 1
    for code, count in sorted(flag_counts.items(), key=lambda x: -x[1]):
        summary.cell(row=row, column=1, value=code)
        summary.cell(row=row, column=2, value=count)
        row += 1

    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 16

    # --- Tier tabs ---
    for tier, label in [
        ("AUTO_APPROVED", "Auto-Approved"),
        ("NEEDS_REVIEW", "Needs Review"),
        ("REJECTED", "Rejected"),
    ]:
        ws = wb.create_sheet(label)
        _write_sheet(ws, tiers[tier])

    wb.save(output_path)
    return tiers
